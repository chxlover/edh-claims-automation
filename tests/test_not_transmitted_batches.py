from __future__ import annotations

import sys
import unittest
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory


PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from core.not_transmitted_batches import (  # noqa: E402
    ClaimBatchWorkbookService,
    ClaimFolderRecord,
    ExtractionResult,
    InvalidClaimRecord,
    LEGACY_WORKBOOK_FORMATS,
    NotTransmittedClaimsRepository,
    WORKBOOK_FORMAT,
)


class FakeCursor:
    def __init__(self, rows):
        self.rows = rows
        self.sql = ""
        self.params = ()

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, traceback):
        return False

    def execute(self, sql, params):
        self.sql = sql
        self.params = params

    def fetchall(self):
        return list(self.rows)


class FakeConnection:
    def __init__(self, rows):
        self.fake_cursor = FakeCursor(rows)
        self.closed = False

    def cursor(self):
        return self.fake_cursor

    def close(self):
        self.closed = True


class AbtcCursor(FakeCursor):
    def execute(self, sql, params):
        self.sql = sql
        self.params = params
        normalized = " ".join(sql.upper().split())
        if "FROM HPATCON " in normalized:
            self.rows = [
                {
                    **claim_row(
                        enccode="ABTC-ELIGIBLE",
                        missing_professional=1,
                    ),
                    "consentdate": None,
                    "authsigndate": None,
                },
                {
                    **claim_row(enccode="ABTC-COMPLETE"),
                    "consentdate": date(2026, 6, 1),
                    "authsigndate": date(2026, 6, 1),
                },
                {
                    **claim_row(enccode="ABTC-MAPPED"),
                    "consentdate": None,
                    "authsigndate": None,
                },
                {
                    **claim_row(enccode="ABTC-NO-PROF"),
                    "consentdate": None,
                    "authsigndate": None,
                },
            ]
        elif "FROM HPHICCLAIMMAP" in normalized:
            self.rows = [{"enccode": "ABTC-MAPPED"}]
        elif "FROM HPROFSERV" in normalized:
            self.rows = [
                {
                    "enccode": "ABTC-ELIGIBLE",
                    "row_count": 1,
                    "has_missing": 1,
                },
                {
                    "enccode": "ABTC-COMPLETE",
                    "row_count": 1,
                    "has_missing": 0,
                },
                {
                    "enccode": "ABTC-MAPPED",
                    "row_count": 1,
                    "has_missing": 1,
                },
            ]
        else:
            self.rows = []


class AbtcConnection(FakeConnection):
    def __init__(self):
        self.fake_cursor = AbtcCursor([])
        self.closed = False


def claim_row(
    *,
    enccode: str,
    missing_professional: int = 0,
    missing_consent: int = 0,
    missing_authorization: int = 0,
    has_profserv: int = 1,
    has_hpatcon1: int = 1,
):
    return {
        "enccode": enccode,
        "hpercode": "000000000012345",
        "admdate": date(2026, 6, 1),
        "disdate": date(2026, 6, 3),
        "patlast": "DELA CRUZ",
        "patfirst": "JUAN",
        "patsuffix": "",
        "patmiddle": "SANTOS",
        "has_profserv": has_profserv,
        "has_hpatcon1": has_hpatcon1,
        "missing_professional_date": missing_professional,
        "missing_consent_date": missing_consent,
        "missing_authorization_date": missing_authorization,
    }


class NotTransmittedRepositoryTests(unittest.TestCase):
    def fetch_rows(self, rows):
        connection = FakeConnection(rows)
        repository = NotTransmittedClaimsRepository(lambda: connection)
        result = repository.fetch(2026)
        return result, connection

    def test_any_one_missing_date_is_included(self):
        for field in (
            "missing_professional",
            "missing_consent",
            "missing_authorization",
        ):
            with self.subTest(field=field):
                values = {field: 1}
                result, _connection = self.fetch_rows(
                    [claim_row(enccode=field, **values)]
                )
                self.assertEqual(len(result.records), 1)
                self.assertEqual(result.records[0].enccode, field)

    def test_all_dates_present_is_excluded(self):
        result, _connection = self.fetch_rows([claim_row(enccode="complete")])

        self.assertEqual(result.records, ())
        self.assertEqual(result.invalid_records, ())

    def test_missing_related_record_goes_to_needs_review(self):
        result, _connection = self.fetch_rows(
            [
                claim_row(
                    enccode="no-prof",
                    has_profserv=0,
                    missing_consent=1,
                ),
                claim_row(
                    enccode="no-consent",
                    has_hpatcon1=0,
                    missing_professional=1,
                ),
            ]
        )

        self.assertEqual(result.records, ())
        self.assertEqual(len(result.invalid_records), 2)
        self.assertIn("Missing hprofserv", result.invalid_records[0].reason)
        self.assertIn("Missing hpatcon1", result.invalid_records[1].reason)

    def test_duplicate_encounter_rows_are_returned_once(self):
        row = claim_row(enccode="same", missing_professional=1)
        result, _connection = self.fetch_rows([row, dict(row)])

        self.assertEqual(len(result.records), 1)

    def test_query_uses_exists_without_date_table_joins(self):
        result, connection = self.fetch_rows(
            [claim_row(enccode="query", missing_consent=1)]
        )
        sql = " ".join(connection.fake_cursor.sql.upper().split())

        self.assertEqual(len(result.records), 1)
        self.assertIn("EXISTS ( SELECT 1 FROM HPROFSERV", sql)
        self.assertIn("EXISTS ( SELECT 1 FROM HPATCON1", sql)
        self.assertNotIn("JOIN HPROFSERV", sql)
        self.assertNotIn("JOIN HPATCON1", sql)
        self.assertIn("LEFT JOIN HPHICCLAIMMAP", sql)
        self.assertIn("HPHICCLAIMMAP.ENCCODE IS NULL", sql)
        self.assertEqual(
            connection.fake_cursor.params,
            ("2026-01-01", "2027-01-01"),
        )

    def test_abtc_excludes_mapped_and_routes_missing_prof_to_review(self):
        repository = NotTransmittedClaimsRepository(AbtcConnection)

        result = repository.fetch(2026, "ABTC")

        self.assertEqual(
            tuple(record.enccode for record in result.records),
            ("ABTC-ELIGIBLE",),
        )
        self.assertEqual(len(result.invalid_records), 1)
        self.assertIn("Missing hprofserv", result.invalid_records[0].reason)


class WorkbookBatchingTests(unittest.TestCase):
    def test_selected_batch_size_and_needs_review_sheet(self):
        records = tuple(
            ClaimFolderRecord(
                f"ENC-{index}",
                f"PATIENT {index:02d}",
                f"{index:015d}",
                date(2026, 6, index),
                date(2026, 6, index + 1),
            )
            for index in range(1, 8)
        )
        invalid = (
            InvalidClaimRecord(
                "REVIEW PATIENT",
                "000000000000099",
                "2026-06-01",
                "2026-06-02",
                "Missing hprofserv record",
            ),
        )
        service = ClaimBatchWorkbookService()

        with TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "batches.xlsx"
            summary = service.export(
                ExtractionResult(records, invalid),
                year=2026,
                destination=workbook_path,
                batch_size=3,
            )

            self.assertEqual(summary.patient_count, 7)
            self.assertEqual(summary.batch_count, 3)
            self.assertEqual(summary.invalid_count, 1)
            self.assertEqual(
                service.list_batch_sheets(workbook_path),
                ("Batch_001", "Batch_002", "Batch_003"),
            )
            self.assertEqual(len(service.read_batch_sheet(workbook_path, "Batch_001")), 3)
            self.assertEqual(len(service.read_batch_sheet(workbook_path, "Batch_002")), 3)
            self.assertEqual(len(service.read_batch_sheet(workbook_path, "Batch_003")), 1)

    def test_new_and_legacy_workbook_formats_are_supported(self):
        from openpyxl import load_workbook

        record = ClaimFolderRecord(
            "ENC-1",
            "PATIENT ONE",
            "000000000000001",
            date(2026, 6, 1),
            date(2026, 6, 2),
        )
        service = ClaimBatchWorkbookService()

        with TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "missing_dates.xlsx"
            service.export(
                ExtractionResult((record,), ()),
                year=2026,
                destination=workbook_path,
            )
            workbook = load_workbook(workbook_path)
            metadata = workbook["_EDH_META"]
            self.assertEqual(metadata["B1"].value, WORKBOOK_FORMAT)
            workbook.close()
            self.assertEqual(
                service.list_batch_sheets(workbook_path),
                ("Batch_001",),
            )

            legacy_path = Path(directory) / "legacy.xlsx"
            workbook = load_workbook(workbook_path)
            workbook["_EDH_META"]["B1"] = next(iter(LEGACY_WORKBOOK_FORMATS))
            workbook.save(legacy_path)
            workbook.close()
            self.assertEqual(
                service.list_batch_sheets(legacy_path),
                ("Batch_001",),
            )

    def test_abtc_workbook_profile_uses_admission_date(self):
        record = ClaimFolderRecord(
            "ABTC-1",
            "ABTC PATIENT",
            "000000000000002",
            date(2026, 6, 1),
            date(2026, 6, 1),
        )
        service = ClaimBatchWorkbookService()

        with TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "abtc.xlsx"
            service.export(
                ExtractionResult((record,), ()),
                year=2026,
                destination=workbook_path,
                claim_type="ABTC",
            )

            self.assertEqual(
                service.workbook_profile(workbook_path),
                ("ABTC", "ADMISSION"),
            )


if __name__ == "__main__":
    unittest.main()
