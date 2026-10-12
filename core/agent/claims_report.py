"""Consolidated final Excel report for the Claims Agent pipeline.

Two sheets: per-claim detail (all section-8 fields, real kind labels
CF4/CF5/eSOA, NOT_RUN for stages never executed) and batch summary
(no double-counting: each claim contributes to exactly one outcome).

Conventions reused: timestamped filenames under ``reports/`` (never
silently overwritten), openpyxl styling after ``fees_checker.write_xlsx``
(bold header, frozen panes, auto filters, green/red outcome fills).

Standalone test:
    python core/agent/claims_report.py
"""

from __future__ import annotations

import os
import sys
import time
from pathlib import Path

if __name__ == "__main__":
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
OK_FILL = PatternFill("solid", fgColor="C6EFCE")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
HOLD_FILL = PatternFill("solid", fgColor="FFE699")

DETAIL_HEADERS = [
    "Batch ID",
    "Claim identifier",
    "Input PDF",
    "Output folder",
    "Claim Processor result",
    "Fees Checker 1st pass",
    "Date Fill result",
    "Final Bill result",
    "Fees Checker 2nd pass",
    "Fees-2 attempt count",
    "XML generation result",
    "CF4 validation",
    "CF5 validation",
    "eSOA validation",
    "Claims Checker status",
    "HBSys Add Claims result",
    "HBSys upload result",
    "Claim Attachments result",
    "Attachment Verification result",
    "Last completed stage",
    "Final outcome",
    "Failure / hold reason",
    "Required manual action",
    "Retry count",
    "Start time",
    "Completion / last-update time",
]

SUMMARY_LABELS = [
    ("total_claims", "Total claims detected"),
    ("completed", "COMPLETED"),
    ("incomplete", "INCOMPLETE"),
    ("held_for_review", "HELD_FOR_REVIEW"),
    ("failed", "FAILED"),
    ("pending", "Interrupted / unfinished"),
    ("xml_generation_failures", "XML generation failures"),
    ("xml_validation_failures", "XML validation failures"),
    ("attachment_failures", "Attachment failures / uncertain"),
    ("needs_manual_action", "Claims requiring manual action"),
]

NOT_RUN = "NOT_RUN"


def _stage_text(record, stage: str) -> str:
    entry = (record.stage_status or {}).get(stage)
    if not entry:
        return NOT_RUN
    status = entry.get("status", "")
    detail = entry.get("detail", "")
    text = status if not detail else f"{status}: {detail}"
    return text[:200]


def _xml_text(record, kind: str) -> str:
    files = (record.xml_result or {}).get("files", {})
    entry = files.get(kind)
    if not entry:
        return NOT_RUN
    if entry.get("valid"):
        series = ",".join(entry.get("series", ())) or "?"
        return f"VALID (series {series})"
    bits = []
    if not entry.get("exists"):
        bits.append("missing")
    elif not entry.get("non_empty"):
        bits.append("empty")
    if entry.get("claim_match") == "MISMATCH":
        bits.append("series mismatch")
    if entry.get("error"):
        bits.append(entry["error"][:120])
    return "INVALID: " + ("; ".join(bits) or "failed checks")


def detail_row(batch_id: str, record) -> list:
    fees2 = (record.stage_status or {}).get("FEES_CHECK_2", {})
    hbsys = record.hbsys or {}
    verification = record.verification or {}
    if verification:
        verify_text = (
            f"{'VERIFIED' if verification.get('verified') else 'NOT VERIFIED'}: "
            f"{verification.get('detail', '')}"[:200]
        )
    else:
        verify_text = NOT_RUN
    return [
        batch_id,
        record.claim_key,
        record.input_pdf or NOT_RUN,
        record.output_path,
        _stage_text(record, "CLAIM_PROCESSING"),
        _stage_text(record, "FEES_CHECK_1"),
        _stage_text(record, "DATE_FILL"),
        _stage_text(record, "FINAL_BILL"),
        fees2.get("status", NOT_RUN)
        if isinstance(fees2, dict) else NOT_RUN,
        fees2.get("evaluations", "") if isinstance(fees2, dict) else "",
        _stage_text(record, "XML_GENERATION"),
        _xml_text(record, "CF4"),
        _xml_text(record, "CF5"),
        _xml_text(record, "eSOA"),
        record.checker_status or NOT_RUN,
        _stage_text(record, "HBSYS_ADD_CLAIMS"),
        (hbsys.get("add_claims_finalize", {}).get("status", NOT_RUN)
         if isinstance(hbsys.get("add_claims_finalize"), dict) else NOT_RUN),
        _stage_text(record, "ATTACHMENTS"),
        verify_text,
        record.last_completed_stage or NOT_RUN,
        record.outcome,
        record.hold_reason or record.last_error,
        record.manual_action,
        record.retries,
        record.started_at,
        record.updated_at,
    ]


def _style_header(sheet, width: int) -> None:
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def write_batch_report(config, batch_id: str, records, summary: dict) -> Path:
    """Write the consolidated report. Never overwrites a prior file."""
    report_dir = Path(config.report_dir)
    report_dir.mkdir(parents=True, exist_ok=True)
    stamp = time.strftime("%Y%m%d_%H%M%S")
    base = f"claims_agent_report_{batch_id}_{stamp}"
    path = report_dir / f"{base}.xlsx"
    suffix = 2
    while path.exists():
        path = report_dir / f"{base}_{suffix}.xlsx"
        suffix += 1

    book = Workbook()
    detail = book.active
    detail.title = "Claims"
    detail.append(DETAIL_HEADERS)
    for record in records:
        detail.append(detail_row(batch_id, record))
    _style_header(detail, len(DETAIL_HEADERS))
    for column in detail.columns:
        longest = max(
            (len(str(cell.value or "")) for cell in column), default=10
        )
        detail.column_dimensions[column[0].column_letter].width = min(longest + 2, 60)

    overview = book.create_sheet("Summary")
    overview.append(["Metric", "Count"])
    for key, label in SUMMARY_LABELS:
        overview.append([label, summary.get(key, 0)])
    _style_header(overview, 2)
    overview.column_dimensions["A"].width = 42
    overview.column_dimensions["B"].width = 14

    # Outcome coloring on the detail sheet (column U = Final outcome).
    fills = {
        "COMPLETED": OK_FILL,
        "FAILED": BAD_FILL,
        "INCOMPLETE": HOLD_FILL,
        "HELD_FOR_REVIEW": HOLD_FILL,
    }
    for row in detail.iter_rows(min_row=2, min_col=21, max_col=21):
        for cell in row:
            fill = fills.get(str(cell.value))
            if fill is not None:
                cell.fill = fill

    book.save(path)
    if getattr(config, "auto_open_report", False):
        try:
            os.startfile(path)  # type: ignore[attr-defined]
        except Exception:
            pass
    return path


if __name__ == "__main__":
    from tempfile import TemporaryDirectory
    from types import SimpleNamespace

    from core.agent.pipeline_state import ClaimRecord  # noqa: E402

    with TemporaryDirectory() as tmp:
        cfg = SimpleNamespace(report_dir=Path(tmp), auto_open_report=False)
        rec = ClaimRecord(batch_id="B1", claim_key="K1", outcome="COMPLETED")
        out = write_batch_report(cfg, "B1", [rec], {"total_claims": 1})
        assert out.is_file() and out.suffix == ".xlsx", out
        out2 = write_batch_report(cfg, "B1", [rec], {"total_claims": 1})
        assert out2 != out and out2.is_file(), (out, out2)
    print("claims_report standalone check passed")
