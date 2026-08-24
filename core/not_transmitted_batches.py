"""Read-only HBSys extraction and Excel folder-batch services."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable, Iterable

from core.activity_logger import ActivityLogger, logger
from core.hbsys_connection import create_hbsys_connection
from core.patient_folder_naming import (
    build_patient_folder_name,
    normalize_folder_date,
    normalize_hpercode,
)


WORKBOOK_FORMAT = "EDH_MISSING_DATE_FILL_BATCHES"
LEGACY_WORKBOOK_FORMATS = frozenset({"EDH_NOT_TRANSMITTED_BATCHES"})
SUPPORTED_WORKBOOK_FORMATS = frozenset(
    {WORKBOOK_FORMAT, *LEGACY_WORKBOOK_FORMATS}
)
WORKBOOK_VERSION = "1"
META_SHEET = "_EDH_META"
BATCH_PREFIX = "Batch_"
BATCH_SIZE = 5
MAX_BATCH_SIZE = 50
CLAIM_TYPES = ("REGULAR", "ABTC")
ABTC_ACCREDITATION_NO = "B02023200"
HEADERS = (
    "PATIENT NAME",
    "HPERCODE",
    "ADMISSION DATE",
    "DISCHARGE DATE",
    "FOLDER NAME",
)


class NotTransmittedBatchError(RuntimeError):
    """Raised when extraction or workbook validation cannot continue safely."""


@dataclass(frozen=True, slots=True)
class ClaimFolderRecord:
    enccode: str
    patient_name: str
    hpercode: str
    admission_date: date
    discharge_date: date

    @property
    def folder_name(self) -> str:
        return build_patient_folder_name(
            self.patient_name,
            self.hpercode,
            self.admission_date,
            self.discharge_date,
        )


@dataclass(frozen=True, slots=True)
class InvalidClaimRecord:
    patient_name: str
    hpercode: str
    admission_date: str
    discharge_date: str
    reason: str


@dataclass(frozen=True, slots=True)
class ExtractionResult:
    records: tuple[ClaimFolderRecord, ...]
    invalid_records: tuple[InvalidClaimRecord, ...]


@dataclass(frozen=True, slots=True)
class WorkbookExportResult:
    path: Path
    patient_count: int
    batch_count: int
    invalid_count: int


@dataclass(frozen=True, slots=True)
class FolderCreationResult:
    created: int
    existing: int
    failed: int
    audit_path: Path


class NotTransmittedClaimsRepository:
    """SELECT-only repository for unfinished Regular/ABTC Date Fill records."""

    def __init__(
        self,
        connection_factory: Callable[[], Any] = create_hbsys_connection,
    ) -> None:
        self.connection_factory = connection_factory

    def available_years(self) -> tuple[int, ...]:
        # Avoid a full-table DISTINCT/YEAR scan merely to populate the GUI.
        # The extraction itself remains authoritative and read-only.
        current = datetime.now().year
        return tuple(range(current, current - 6, -1))

    def fetch(self, year: int, claim_type: str = "REGULAR") -> ExtractionResult:
        if int(year) < 2000 or int(year) > 2100:
            raise ValueError("Select a valid claim year")
        normalized_type = _normalize_claim_type(claim_type)
        if normalized_type == "ABTC":
            rows = self._fetch_abtc_rows(int(year))
        else:
            sql = """
            SELECT
                hadmlog.enccode,
                hadmlog.hpercode,
                hadmlog.admdate,
                hadmlog.disdate,
                hperson.patlast,
                hperson.patfirst,
                hperson.patsuffix,
                hperson.patmiddle,
                EXISTS (
                    SELECT 1
                    FROM hprofserv
                    WHERE hprofserv.enccode = hadmlog.enccode
                ) AS has_profserv,
                EXISTS (
                    SELECT 1
                    FROM hpatcon1
                    WHERE hpatcon1.enccode = hadmlog.enccode
                ) AS has_hpatcon1,
                EXISTS (
                    SELECT 1
                    FROM hprofserv
                    WHERE hprofserv.enccode = hadmlog.enccode
                      AND hprofserv.pdoctorsigndate IS NULL
                ) AS missing_professional_date,
                EXISTS (
                    SELECT 1
                    FROM hpatcon1
                    WHERE hpatcon1.enccode = hadmlog.enccode
                      AND hpatcon1.consentdate IS NULL
                ) AS missing_consent_date,
                EXISTS (
                    SELECT 1
                    FROM hpatcon1
                    WHERE hpatcon1.enccode = hadmlog.enccode
                      AND hpatcon1.authsigndate IS NULL
                ) AS missing_authorization_date
            FROM hadmlog
            INNER JOIN henctr
                ON hadmlog.enccode = henctr.enccode
               AND henctr.phicclaim = 'Y'
            INNER JOIN hperson
                ON hadmlog.hpercode = hperson.hpercode
            LEFT JOIN hphicclaimmap
                ON hadmlog.enccode = hphicclaimmap.enccode
            WHERE hadmlog.disdate >= %s
              AND hadmlog.disdate < %s
              AND hphicclaimmap.enccode IS NULL
            ORDER BY hadmlog.disdate ASC,
                     hperson.patlast,
                     hperson.patfirst,
                     hadmlog.hpercode
        """
            rows = self._select(
                sql,
                (f"{int(year):04d}-01-01", f"{int(year) + 1:04d}-01-01"),
            )
        records: list[ClaimFolderRecord] = []
        invalid: list[InvalidClaimRecord] = []
        seen_encounters: set[str] = set()
        for row in rows:
            encounter = str(row.get("enccode") or "").strip()
            if encounter and encounter in seen_encounters:
                continue
            if encounter:
                seen_encounters.add(encounter)
            patient_name = _patient_name(row)
            hpercode = normalize_hpercode(row.get("hpercode"))
            admission = _as_date(row.get("admdate"))
            discharge = _as_date(row.get("disdate"))
            missing_relations: list[str] = []
            if not bool(row.get("has_profserv")):
                missing_relations.append("Missing hprofserv record")
            if not bool(row.get("has_hpatcon1")):
                missing_relations.append("Missing hpatcon1 record")
            if missing_relations:
                invalid.append(
                    InvalidClaimRecord(
                        patient_name,
                        hpercode,
                        str(row.get("admdate") or ""),
                        str(row.get("disdate") or ""),
                        "; ".join(missing_relations),
                    )
                )
                continue
            has_missing_date = any(
                bool(row.get(field))
                for field in (
                    "missing_professional_date",
                    "missing_consent_date",
                    "missing_authorization_date",
                )
            )
            if not has_missing_date:
                continue
            problems: list[str] = []
            if not patient_name:
                problems.append("Missing patient name")
            if len(hpercode) != 15:
                problems.append("Missing or invalid HPERCODE")
            if admission is None:
                problems.append("Missing or invalid admission date")
            if discharge is None:
                problems.append("Missing or invalid discharge date")
            if admission and discharge and admission > discharge:
                problems.append("Admission date is after discharge date")
            if problems:
                invalid.append(
                    InvalidClaimRecord(
                        patient_name,
                        hpercode,
                        str(row.get("admdate") or ""),
                        str(row.get("disdate") or ""),
                        "; ".join(problems),
                    )
                )
                continue
            records.append(
                ClaimFolderRecord(
                    encounter,
                    patient_name,
                    hpercode,
                    admission,
                    discharge,
                )
            )
        records.sort(
            key=lambda item: (
                (
                    item.admission_date
                    if normalized_type == "ABTC"
                    else item.discharge_date
                ),
                item.patient_name,
                item.hpercode,
                item.enccode,
            )
        )
        return ExtractionResult(tuple(records), tuple(invalid))

    def _fetch_abtc_rows(self, year: int) -> list[dict[str, Any]]:
        """Read ABTC rows in small indexed queries to avoid a heavy map join."""
        connection = None
        try:
            connection = self.connection_factory()
            with connection.cursor() as cursor:
                cursor.execute(
                    """
                    SELECT
                        hpatcon.enccode,
                        hpatcon.hpercode,
                        hpatcon1.padmissiondatetime AS admdate,
                        hpatcon1.pdischargedatetime AS disdate,
                        hpatcon1.consentdate,
                        hpatcon1.authsigndate,
                        hperson.patlast,
                        hperson.patfirst,
                        hperson.patsuffix,
                        hperson.patmiddle
                    FROM hpatcon
                    INNER JOIN hpatcon1
                        ON hpatcon1.enccode = hpatcon.enccode
                    INNER JOIN hperson
                        ON hperson.hpercode = hpatcon.hpercode
                    WHERE hpatcon.accreno = %s
                      AND hpatcon1.padmissiondatetime >= %s
                      AND hpatcon1.padmissiondatetime < %s
                    ORDER BY hpatcon1.padmissiondatetime,
                             hperson.patlast,
                             hperson.patfirst,
                             hpatcon.hpercode
                    """,
                    (
                        ABTC_ACCREDITATION_NO,
                        f"{year:04d}-01-01",
                        f"{year + 1:04d}-01-01",
                    ),
                )
                base_rows = list(cursor.fetchall())
                codes = tuple(
                    str(row.get("enccode") or "").strip()
                    for row in base_rows
                    if str(row.get("enccode") or "").strip()
                )
                mapped: set[str] = set()
                professional: dict[str, tuple[bool, bool]] = {}
                for chunk in _chunks(codes, 400):
                    placeholders = ",".join("%s" for _ in chunk)
                    cursor.execute(
                        "SELECT DISTINCT enccode FROM hphicclaimmap "
                        f"WHERE enccode IN ({placeholders})",
                        chunk,
                    )
                    mapped.update(
                        str(row.get("enccode") or "").strip()
                        for row in cursor.fetchall()
                    )
                    cursor.execute(
                        """
                        SELECT enccode,
                               COUNT(*) AS row_count,
                               MAX(pdoctorsigndate IS NULL) AS has_missing
                        FROM hprofserv
                        WHERE enccode IN ("""
                        + placeholders
                        + ") GROUP BY enccode",
                        chunk,
                    )
                    for row in cursor.fetchall():
                        professional[str(row.get("enccode") or "").strip()] = (
                            bool(row.get("row_count")),
                            bool(row.get("has_missing")),
                        )
        except Exception as exc:
            raise NotTransmittedBatchError(
                "Unable to read ABTC patients with missing Date Fill fields "
                "from HBSys"
            ) from exc
        finally:
            if connection is not None:
                connection.close()

        rows: list[dict[str, Any]] = []
        for source in base_rows:
            encounter = str(source.get("enccode") or "").strip()
            if not encounter or encounter in mapped:
                continue
            has_profserv, missing_professional = professional.get(
                encounter, (False, False)
            )
            row = dict(source)
            row.update(
                {
                    "has_profserv": has_profserv,
                    "has_hpatcon1": True,
                    "missing_professional_date": missing_professional,
                    "missing_consent_date": source.get("consentdate") is None,
                    "missing_authorization_date": (
                        source.get("authsigndate") is None
                    ),
                }
            )
            rows.append(row)
        return rows

    def _select(self, sql: str, params: tuple[Any, ...]) -> list[dict[str, Any]]:
        connection = None
        try:
            connection = self.connection_factory()
            with connection.cursor() as cursor:
                cursor.execute(sql, params)
                return list(cursor.fetchall())
        except Exception as exc:
            raise NotTransmittedBatchError(
                "Unable to read Regular patients with missing Date Fill fields "
                "from HBSys"
            ) from exc
        finally:
            if connection is not None:
                connection.close()


class ClaimBatchWorkbookService:
    """Create, validate, and consume EDH folder-batch workbooks."""

    def __init__(self, activity_logger: ActivityLogger = logger) -> None:
        self.logger = activity_logger

    def export(
        self,
        result: ExtractionResult,
        *,
        year: int,
        destination: str | Path,
        batch_size: int = BATCH_SIZE,
        claim_type: str = "REGULAR",
    ) -> WorkbookExportResult:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError as exc:
            raise NotTransmittedBatchError(
                "openpyxl is required. Run install_requirements.py first."
            ) from exc
        normalized_type = _normalize_claim_type(claim_type)
        if not result.records and not result.invalid_records:
            raise NotTransmittedBatchError(
                f"No eligible {normalized_type} patients with missing Date Fill fields "
                f"were found for {year}. Transmitted claims are excluded."
            )
        batch_size = int(batch_size)
        if not 1 <= batch_size <= MAX_BATCH_SIZE:
            raise ValueError(
                f"Patients per sheet must be from 1 to {MAX_BATCH_SIZE}"
            )
        path = Path(destination)
        if path.suffix.lower() != ".xlsx":
            path = path.with_suffix(".xlsx")
        path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        metadata = workbook.active
        metadata.title = META_SHEET
        metadata.append(("FORMAT", WORKBOOK_FORMAT))
        metadata.append(("VERSION", WORKBOOK_VERSION))
        metadata.append(("YEAR", int(year)))
        metadata.append(("EXPORTED_AT", datetime.now().isoformat(timespec="seconds")))
        metadata.append(("PATIENT_COUNT", len(result.records)))
        metadata.append(("BATCH_SIZE", batch_size))
        metadata.append(("CLAIM_TYPE", normalized_type))
        metadata.append(
            ("DATE_BASIS", "ADMISSION" if normalized_type == "ABTC" else "DISCHARGE")
        )
        metadata.sheet_state = "hidden"

        header_fill = PatternFill("solid", fgColor="2F6FED")
        header_font = Font(color="FFFFFF", bold=True)
        records = result.records
        for offset in range(0, len(records), batch_size):
            batch_number = (offset // batch_size) + 1
            sheet = workbook.create_sheet(f"{BATCH_PREFIX}{batch_number:03d}")
            sheet.append(HEADERS)
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
            for record in records[offset : offset + batch_size]:
                sheet.append(
                    (
                        record.patient_name,
                        record.hpercode,
                        record.admission_date,
                        record.discharge_date,
                        record.folder_name,
                    )
                )
            sheet.freeze_panes = "A2"
            sheet.column_dimensions["A"].width = 38
            sheet.column_dimensions["B"].width = 19
            sheet.column_dimensions["C"].width = 18
            sheet.column_dimensions["D"].width = 18
            sheet.column_dimensions["E"].width = 78
            for row in sheet.iter_rows(min_row=2, min_col=3, max_col=4):
                for cell in row:
                    cell.number_format = "yyyy-mm-dd"

        if result.invalid_records:
            sheet = workbook.create_sheet("Needs_Review")
            sheet.append((*HEADERS[:4], "VALIDATION ISSUE"))
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="DC2626")
                cell.font = header_font
            for item in result.invalid_records:
                sheet.append(
                    (
                        item.patient_name,
                        item.hpercode,
                        item.admission_date,
                        item.discharge_date,
                        item.reason,
                    )
                )

        workbook.save(path)
        batch_count = (len(records) + batch_size - 1) // batch_size
        self.logger.success(
            f"Missing Date Fill workbook exported: {path} "
            f"patients={len(records)} batches={batch_count}"
        )
        return WorkbookExportResult(
            path,
            len(records),
            batch_count,
            len(result.invalid_records),
        )

    def list_batch_sheets(self, workbook_path: str | Path) -> tuple[str, ...]:
        workbook = self._load_validated_workbook(workbook_path, read_only=True)
        try:
            return tuple(
                name
                for name in workbook.sheetnames
                if name.startswith(BATCH_PREFIX)
            )
        finally:
            workbook.close()

    def workbook_profile(self, workbook_path: str | Path) -> tuple[str, str]:
        """Return claim type/date basis, defaulting legacy workbooks safely."""
        workbook = self._load_validated_workbook(workbook_path, read_only=True)
        try:
            metadata = self._read_metadata(workbook)
            claim_type = _normalize_claim_type(
                metadata.get("CLAIM_TYPE", "REGULAR")
            )
            default_basis = "ADMISSION" if claim_type == "ABTC" else "DISCHARGE"
            date_basis = str(metadata.get("DATE_BASIS", default_basis)).upper()
            if date_basis not in {"ADMISSION", "DISCHARGE"}:
                raise NotTransmittedBatchError(
                    "Workbook has an invalid Date Fill basis"
                )
            if claim_type == "ABTC" and date_basis != "ADMISSION":
                raise NotTransmittedBatchError(
                    "ABTC workbook must use the admission date"
                )
            if claim_type == "REGULAR" and date_basis != "DISCHARGE":
                raise NotTransmittedBatchError(
                    "Regular workbook must use the discharge date"
                )
            return claim_type, date_basis
        finally:
            workbook.close()

    def read_batch_sheet(
        self,
        workbook_path: str | Path,
        sheet_name: str,
    ) -> tuple[ClaimFolderRecord, ...]:
        workbook = self._load_validated_workbook(workbook_path, read_only=True)
        try:
            if sheet_name not in workbook.sheetnames or not sheet_name.startswith(BATCH_PREFIX):
                raise NotTransmittedBatchError("Select a valid Batch sheet")
            sheet = workbook[sheet_name]
            metadata = self._read_metadata(workbook)
            try:
                batch_limit = int(metadata.get("BATCH_SIZE", BATCH_SIZE))
            except (TypeError, ValueError) as exc:
                raise NotTransmittedBatchError(
                    "Workbook has an invalid Patients per sheet value"
                ) from exc
            if not 1 <= batch_limit <= MAX_BATCH_SIZE:
                raise NotTransmittedBatchError(
                    "Workbook Patients per sheet value is outside the safe limit"
                )
            headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
            if headers != HEADERS:
                raise NotTransmittedBatchError(
                    f"Workbook sheet {sheet_name} has invalid or changed headers"
                )
            records: list[ClaimFolderRecord] = []
            for values in sheet.iter_rows(min_row=2, values_only=True):
                if not any(value not in (None, "") for value in values):
                    continue
                patient_name = " ".join(str(values[0] or "").strip().split())
                hpercode = normalize_hpercode(values[1])
                admission = _as_date(values[2])
                discharge = _as_date(values[3])
                stored_folder = str(values[4] or "").strip()
                if not patient_name or len(hpercode) != 15 or not admission or not discharge:
                    raise NotTransmittedBatchError(
                        f"{sheet_name} contains an incomplete patient row"
                    )
                if admission > discharge:
                    raise NotTransmittedBatchError(
                        f"{sheet_name} contains an invalid confinement period"
                    )
                record = ClaimFolderRecord(
                    "",
                    patient_name,
                    hpercode,
                    admission,
                    discharge,
                )
                if stored_folder != record.folder_name:
                    raise NotTransmittedBatchError(
                        f"Folder Name was changed or is invalid for {patient_name}"
                    )
                records.append(record)
            if not records:
                raise NotTransmittedBatchError(f"{sheet_name} has no patient rows")
            if len(records) > batch_limit:
                raise NotTransmittedBatchError(
                    f"{sheet_name} exceeds its {batch_limit}-patient sheet limit"
                )
            return tuple(records)
        finally:
            workbook.close()

    def create_folders(
        self,
        workbook_path: str | Path,
        sheet_name: str,
        *,
        output_folder: str | Path,
        reports_folder: str | Path,
    ) -> FolderCreationResult:
        records = self.read_batch_sheet(workbook_path, sheet_name)
        claim_type, date_basis = self.workbook_profile(workbook_path)
        output_root = Path(output_folder).resolve()
        output_root.mkdir(parents=True, exist_ok=True)
        audit_root = Path(reports_folder)
        audit_root.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        audit_path = audit_root / f"missing_date_fill_folder_creation_{timestamp}.csv"
        audit_rows: list[tuple[str, ...]] = []
        created = existing = failed = 0
        for record in records:
            destination = (output_root / record.folder_name).resolve()
            if destination.parent != output_root:
                action = "FAILED_UNSAFE_PATH"
                failed += 1
            elif destination.exists() and destination.is_dir():
                action = "ALREADY_EXISTS"
                existing += 1
            elif destination.exists():
                action = "FAILED_PATH_IS_FILE"
                failed += 1
            else:
                try:
                    destination.mkdir()
                    action = "CREATED"
                    created += 1
                except Exception as exc:
                    action = f"FAILED: {exc}"
                    failed += 1
            audit_rows.append(
                (
                    datetime.now().isoformat(timespec="seconds"),
                    str(Path(workbook_path)),
                    sheet_name,
                    claim_type,
                    date_basis,
                    record.patient_name,
                    record.hpercode,
                    record.admission_date.isoformat(),
                    record.discharge_date.isoformat(),
                    str(destination),
                    action,
                )
            )
        with audit_path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle)
            writer.writerow(
                (
                    "timestamp",
                    "workbook",
                    "sheet",
                    "claim_type",
                    "date_basis",
                    "patient_name",
                    "hpercode",
                    "admission_date",
                    "discharge_date",
                    "folder_path",
                    "action",
                )
            )
            writer.writerows(audit_rows)
        self.logger.success(
            f"Missing Date Fill folders complete: sheet={sheet_name} "
            f"created={created} existing={existing} failed={failed}"
        )
        return FolderCreationResult(created, existing, failed, audit_path)

    @staticmethod
    def _load_validated_workbook(workbook_path: str | Path, *, read_only: bool):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise NotTransmittedBatchError(
                "openpyxl is required. Run install_requirements.py first."
            ) from exc
        path = Path(workbook_path)
        if not path.is_file() or path.suffix.lower() != ".xlsx":
            raise NotTransmittedBatchError("Select a valid .xlsx batch workbook")
        try:
            workbook = load_workbook(path, read_only=read_only, data_only=True)
        except Exception as exc:
            raise NotTransmittedBatchError("Unable to open the selected workbook") from exc
        if META_SHEET not in workbook.sheetnames:
            workbook.close()
            raise NotTransmittedBatchError("This is not an EDH folder-batch workbook")
        metadata = ClaimBatchWorkbookService._read_metadata(workbook)
        if (
            metadata.get("FORMAT") not in SUPPORTED_WORKBOOK_FORMATS
            or metadata.get("VERSION") != WORKBOOK_VERSION
        ):
            workbook.close()
            raise NotTransmittedBatchError("Unsupported or invalid workbook format")
        return workbook

    @staticmethod
    def _read_metadata(workbook) -> dict[str, str]:
        return {
            str(row[0].value or ""): str(row[1].value or "")
            for row in workbook[META_SHEET].iter_rows(min_col=1, max_col=2)
        }


def _patient_name(row: dict[str, Any]) -> str:
    last = str(row.get("patlast") or "").strip().upper()
    given = " ".join(
        str(row.get(field) or "").strip().upper()
        for field in ("patfirst", "patsuffix", "patmiddle")
        if str(row.get(field) or "").strip()
    )
    return " ".join(f"{last}, {given}".strip(" ,").split())


def _normalize_claim_type(value: object) -> str:
    normalized = str(value or "REGULAR").strip().upper()
    if normalized not in CLAIM_TYPES:
        raise ValueError("Claim Type must be Regular or ABTC")
    return normalized


def _chunks(values: tuple[str, ...], size: int) -> Iterable[tuple[str, ...]]:
    for offset in range(0, len(values), size):
        yield values[offset : offset + size]


def _as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    normalized = normalize_folder_date(value)
    if not normalized:
        return None
    try:
        return datetime.strptime(normalized, "%Y%m%d").date()
    except ValueError:
        return None


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Export Missing Date Fill batches")
    parser.add_argument("--year", type=int, default=datetime.now().year)
    parser.add_argument(
        "--claim-type",
        choices=CLAIM_TYPES,
        default="REGULAR",
    )
    parser.add_argument("--output", type=Path)
    arguments = parser.parse_args()
    target = arguments.output or (
        Path(__file__).resolve().parent.parent
        / "reports"
        / (
            f"missing_date_fill_{arguments.claim_type.lower()}_"
            f"{arguments.year}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        )
    )
    extraction = NotTransmittedClaimsRepository().fetch(
        arguments.year,
        arguments.claim_type,
    )
    summary = ClaimBatchWorkbookService().export(
        extraction,
        year=arguments.year,
        destination=target,
        claim_type=arguments.claim_type,
    )
    print(summary)
